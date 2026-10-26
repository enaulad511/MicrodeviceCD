# -*- coding: utf-8 -*-
"""Lectura/escritura de los libros .xlsx con los que la app guarda sus datos.

Cada guardado (corrida PCR, corrida EmStat, export de análisis, pestañas
manuales) es UN archivo con varias hojas, para poder moverlo de un lado a otro
sin perder hermanos (docs/almacenamiento_por_experimento.md).

- ``write_workbook`` escribe a ``<path>.tmp`` y renombra: un fallo a media
  escritura nunca deja un .xlsx corrupto con el nombre final.
- ``read_sheet_text`` devuelve las filas como texto, igual que ``csv.reader``
  (celda vacía → ``""``, sin celdas vacías al final), así los loaders conservan
  su lógica de parseo de cuando leían CSV.

``openpyxl`` se importa perezoso dentro de cada función: un Pi sin la librería
arranca igual y el fallo aparece solo al guardar/cargar (donde se maneja).
"""

import json
import math
import os

__author__ = "Edisson A. Naula"
__date__ = "$ 26/10/2026 at 10:00 a.m. $"


XLSX_FILETYPES = [("Excel workbook", "*.xlsx"), ("All files", "*.*")]


def xlsx_path(path: str) -> str:
    """Normaliza la extensión elegida en un diálogo "Save as" a ``.xlsx``."""
    base, ext = os.path.splitext(path)
    return path if ext.lower() == ".xlsx" else f"{base if ext.lower() == '.csv' else path}.xlsx"


def _cell(value):
    # openpyxl acepta números, texto, bool y None; el resto (numpy, listas, dicts)
    # se convierte para que una celda rara no tumbe el guardado completo.
    try:
        import numpy as np

        if isinstance(value, np.generic):
            value = value.item()
    except ImportError:
        pass
    if isinstance(value, float) and not math.isfinite(value):
        # NaN/inf como número dejan un libro que Excel reporta corrupto; como texto
        # ("nan") Excel lo abre y float("nan") lo relee igual que en el CSV de antes.
        return repr(value)
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (list, dict, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


_exact_floats_done = False


def _exact_floats() -> None:
    """openpyxl escribe los números con ``"%.16g"`` y pierde el 17.º dígito (~1e-14
    en 90 °C). Parece nada, pero el crudo del termopar se reconstruye invirtiendo la
    EMA y cae en múltiplos exactos de 0.25 °C: con 16 dígitos los saltos de justo
    3.0 °C cruzan un umbral ``> 3`` de un lado a otro (medido en Newn1 21/10). ``repr``
    es el formato más corto que relee el float exacto, como hacía el CSV. Si otra
    versión de openpyxl cambia este interno, se queda en 16 dígitos sin romper nada.
    """
    global _exact_floats_done
    if _exact_floats_done:
        return
    _exact_floats_done = True
    try:
        from openpyxl.cell import _writer

        original = _writer.safe_string
    except (ImportError, AttributeError):
        return

    def safe_string(value):
        if type(value) is float and math.isfinite(value):
            return repr(value)
        return original(value)

    _writer.safe_string = safe_string


def write_workbook(path: str, sheets: dict) -> None:
    """Escribe ``{nombre_hoja: filas}`` (filas = iterables de celdas; la primera
    suele ser el encabezado). El orden del dict es el orden de las hojas.

    Escritura atómica: ``.tmp`` + ``os.replace``. Lanza la excepción original si
    algo falla (el llamador decide si hay respaldo o solo mensaje).
    """
    from openpyxl import Workbook

    _exact_floats()
    wb = Workbook(write_only=True)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=str(name)[:31])
        for row in rows:
            ws.append([_cell(v) for v in row])
    tmp = f"{path}.tmp"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def sheet_names(path: str) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _as_text(value) -> str:
    if value is None:
        return ""
    # repr de float = round-trip exacto con float(); 3.0 de un entero guardado
    # como float sigue parseando con int(float(x)).
    return repr(value) if isinstance(value, float) else str(value)


def read_sheets_text(path: str) -> dict[str, list[list[str]]]:
    """Todas las hojas como ``{nombre: filas de texto}`` (ver ``read_sheet_text``)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out = {}
        for ws in wb.worksheets:
            # getattr: una chartsheet (libro editado a mano) no tiene celdas; además
            # los stubs tipan ReadOnlyWorksheet.values como método de Worksheet.
            values = getattr(ws, "values", None)
            if values is None:
                continue
            rows = []
            for raw in values:
                row = [_as_text(v) for v in raw]
                while row and row[-1] == "":
                    row.pop()
                rows.append(row)
            out[ws.title] = rows
        return out
    finally:
        wb.close()


def read_sheet_text(path: str, name: str) -> list[list[str]] | None:
    """Filas de la hoja ``name`` como texto, o ``None`` si el libro no la tiene."""
    return read_sheets_text(path).get(name)


def num_or_text(value):
    """Texto numérico → número (para tablas que vienen de un Treeview); lo demás igual."""
    if isinstance(value, str):
        try:
            return int(value)
        except ValueError:
            pass
        try:
            return float(value)
        except ValueError:
            return value if value != "" else None
    return value


def kv_rows(pairs) -> list[list]:
    """Hoja ``key | value`` a partir de pares (metadata, settings)."""
    return [["key", "value"]] + [[k, v] for k, v in pairs]


def flatten_settings(data, prefix: str = "") -> list[tuple[str, object]]:
    """Aplana un dict anidado a pares ``"a.b.c" → valor`` (listas como JSON)."""
    pairs = []
    for key, value in data.items():
        full = f"{prefix}{key}"
        if isinstance(value, dict):
            pairs.extend(flatten_settings(value, f"{full}."))
        else:
            pairs.append((full, value))
    return pairs
